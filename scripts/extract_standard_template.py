"""Read the approved Excel template; emit sanitized Sheets API data, never a workbook.

Run with the bundled document Python (openpyxl). No runtime dependency is added.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX


def extract(source: Path, target: Path):
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    book = openpyxl.load_workbook(source, data_only=False)
    sheet = book["양식예시(복사해서 사용)"]
    theme = ElementTree.fromstring(book.loaded_theme)
    scheme = theme.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme")
    theme_colors = [list(item)[0].get("val", list(item)[0].get("lastClr", "000000")) for item in scheme]
    theme_colors[:4] = [theme_colors[i] for i in (1, 0, 3, 2)]

    def color(value):
        if value is None or value.type == "auto":
            return {"red": 0, "green": 0, "blue": 0}
        rgb = (value.rgb if value.type == "rgb" else
               theme_colors[value.theme] if value.type == "theme" else
               COLOR_INDEX[value.indexed] if value.type == "indexed" and value.indexed < 64 else "000000")[-6:]
        channels = [int(rgb[i:i+2], 16)/255 for i in (0, 2, 4)]
        tint = value.tint or 0
        return dict(zip(("red", "green", "blue"),
                        [v * (1+tint) if tint < 0 else v * (1-tint)+tint for v in channels]))

    def formatting(cell):
        text = {"fontFamily": cell.font.name or "Malgun Gothic", "fontSize": round(cell.font.sz or 11),
                "bold": bool(cell.font.b), "italic": bool(cell.font.i), "foregroundColor": color(cell.font.color)}
        result = {"textFormat": text, "verticalAlignment": {"center": "MIDDLE", "top": "TOP", "bottom": "BOTTOM"}.get(cell.alignment.vertical, "BOTTOM"),
                  "wrapStrategy": "WRAP" if cell.alignment.wrap_text else "CLIP"}
        if cell.alignment.horizontal in ("left", "center", "right"):
            result["horizontalAlignment"] = cell.alignment.horizontal.upper()
        if cell.fill.patternType == "solid":
            result["backgroundColor"] = color(cell.fill.fgColor)
        borders = {}
        for edge in ("top", "bottom", "left", "right"):
            side = getattr(cell.border, edge)
            if side and side.style:
                borders[edge] = {"style": {"thin": "SOLID", "medium": "SOLID_MEDIUM", "thick": "SOLID_THICK",
                    "dashed": "DASHED", "dotted": "DOTTED", "double": "DOUBLE"}.get(side.style, "SOLID"),
                    "color": color(side.color)}
        if borders:
            result["borders"] = borders
        if cell.number_format == "@":
            result["numberFormat"] = {"type": "TEXT", "pattern": "@"}
        elif cell.number_format != "General":
            result["numberFormat"] = {"type": "NUMBER", "pattern": cell.number_format}
        return result

    cells = []
    for row in range(1, 21):
        for col in range(1, 19):
            cell = sheet.cell(row, col)
            if row > 3 and col < 16:
                continue
            item = {"userEnteredFormat": formatting(cell)}
            # Headers, formulas in the prototype row, and helper lists only.
            keep = row <= 2 or col >= 16 or row == 3 and col in (4, 11, 12, 13, 14)
            if keep and cell.value is not None:
                key = "formulaValue" if cell.data_type == "f" else "numberValue" if isinstance(cell.value, (int, float)) else "stringValue"
                item["userEnteredValue"] = {key: cell.value}
            cells.append([row-1, col-1, item])
    for validation in sheet.data_validations.dataValidation:
        if validation.type != "list" or not validation.formula1.startswith('"'):
            raise ValueError("Only explicit list validation is supported")
        values = validation.formula1.strip('"').split(",")
        for col in range(1, 19):
            if sheet.cell(3, col).coordinate in validation:
                item = next(item for r, c, item in cells if (r, c) == (2, col-1))
                item["dataValidation"] = {"condition": {"type": "ONE_OF_LIST", "values": [
                    {"userEnteredValue": v} for v in values]}, "strict": False, "showCustomUi": True}
    widths = [64] * 26
    for dimension in sheet.column_dimensions.values():
        for index in range((dimension.min or 1)-1, dimension.max or dimension.min or 1):
            widths[index] = round(dimension.width * 7 + 5)
    result = {"version": 1, "source_sheet": sheet.title, "source_sha256": before,
              "row_count": 1000, "column_widths": widths,
              "row_heights": [round((sheet.row_dimensions[i].height or 17.25)*4/3) for i in (1, 2, 3)],
              "cells": cells}
    book.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Sanitized template created; original workbook unchanged.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("target", type=Path)
    args = parser.parse_args()
    extract(args.source, args.target)
