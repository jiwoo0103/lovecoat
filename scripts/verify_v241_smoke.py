"""Exercise installed scope changes, real AI, restart and Excel append offline."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from coatdesk_v2.storage import Repository  # noqa: E402


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--exe", type=Path, required=True)
    parser.add_argument("--work", type=Path, required=True)
    args = parser.parse_args()
    work = args.work.resolve()
    if not work.is_relative_to(ROOT / ".local"):
        raise ValueError("Verification data must stay in project/.local")
    data, target = work / "installed-data", work / "installed-export.xlsx"
    if data.exists() or target.exists():
        raise FileExistsError("Use fresh verification data and workbook")
    environment = os.environ.copy()
    environment["QT_QPA_PLATFORM"] = "offscreen"
    environment["QT_QUICK_BACKEND"] = "software"
    rounds = []
    preserved_sheets = {}
    for index, (start, end, errors, codes) in enumerate(((1, 6, "5", ["0001", "0002", "0003", "0004"]),
                                                       (2, 4, "", ["0002", "0003", "0004"]),
                                                       (1, 6, "5", ["0001", "0002", "0003", "0004"]))):
        capture = work / f"installed-run-{index+1}.png"
        command = [str(args.exe.resolve()), "--data-dir", str(data), "--smoke-test", str(capture),
                   "--smoke-scope", str(start), str(end), "--smoke-errors", errors, "--smoke-ocr",
                   "--smoke-export", str(target), "--software-rendering"]
        if index == 0:
            command += ["--source", str(work / "photos"), "--range-start", "1", "--range-end", "2"]
        subprocess.run(command, env=environment, check=True, timeout=120, capture_output=True)
        report = json.loads(capture.with_suffix(".json").read_text(encoding="utf-8"))
        assert report["ok"] and report["scope"] == [start, end]
        assert report["empty_photo_codes"] == (["0005"] if errors else [])
        assert report["counts"]["total"] == len(codes) and report["counts"]["confirmed"] == len(codes)
        book = load_workbook(target)
        assert len(book.sheetnames) == index+1
        sheet = book[book.sheetnames[-1]]
        assert [sheet.cell(row, 3).value for row in range(3, 3+len(codes))] == codes
        assert all(sheet.cell(row, 5).value in ("자켓", "코트", "조끼", "기타") for row in range(3, 3+len(codes)))
        assert all(sheet.cell(row, 4).data_type == "f" for row in range(3, 3+len(codes)))
        book.close()
        with zipfile.ZipFile(target) as archive:
            for name, original in preserved_sheets.items():
                assert archive.read(name) == original
            preserved_sheets = {name: archive.read(name) for name in archive.namelist()
                                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")}
        repo = Repository(data)
        workspace = repo.load(repo.setting("last_workspace"))
        assert all(g.suggestion["category"]["thresholds"] == {"sleeve": .9, "length": .9} for g in workspace.active)
        rounds.append({"scope": [start, end], "rows": len(codes), "photo_empty": report["empty_photo_codes"],
                       "warnings": report["warnings"], "workbook_sheets": index+1})
        print(json.dumps(rounds[-1]), flush=True)
    integrity = json.loads((work / "source-integrity.json").read_text(encoding="utf-8"))
    originals_unchanged = all(hashlib.sha256(Path(row["source"]).read_bytes()).hexdigest() == row["sha256"] for row in integrity)
    result = {"ok": originals_unchanged, "rounds": rounds, "originals_unchanged": originals_unchanged,
              "earlier_sheet_xml_unchanged": True, "exe_sha256": hashlib.sha256(args.exe.read_bytes()).hexdigest()}
    (work / "end-to-end-report.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
