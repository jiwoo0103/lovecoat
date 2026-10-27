"""Local Excel export. Existing worksheets and their package parts stay intact."""
from __future__ import annotations

import copy
import ctypes
import hashlib
import io
import json
import os
import re
import sys
import tempfile
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from xml.dom import minidom
from xml.parsers.expat import ExpatError
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PySide6.QtCore import QLockFile

from .domain import Workspace, fingerprint, validate_fields
from .storage import Repository

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
DOC_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CONTENT = "http://schemas.openxmlformats.org/package/2006/content-types"
MODIFIED_PARTS = {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/styles.xml", "[Content_Types].xml"}


@dataclass
class ExportPlan:
    workspace_id: str
    workspace_digest: str
    path: str
    file_digest: str | None
    title: str
    rows: list[list]
    counts: dict
    template_digest: str


def _template():
    root = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1]))
    try:
        data = json.loads((root / "resources/standard_sheet.json").read_text(encoding="utf-8"))
        if (data["version"] != 1 or len(data["row_heights"]) != 3 or
                not 3 <= data["row_count"] <= 1048576 or not data["column_widths"]):
            raise ValueError
        if not all(type(v) is int and v > 0 for v in data["row_heights"] + data["column_widths"]):
            raise ValueError
        for row, col, item in data["cells"]:
            value = item.get("userEnteredValue", {})
            if row >= 2 and col < 15 and value and not (
                    row == 2 and (col == 3 or 10 <= col < 14) and "formulaValue" in value):
                raise ValueError
        return data
    except (OSError, KeyError, TypeError, ValueError):
        raise ValueError("내장 엑셀 양식을 읽을 수 없습니다. 프로그램 설치 파일을 확인하세요.") from None


def next_title(base: str, titles: list[str]) -> str:
    base = base.strip()
    if not base or len(base) > 31 or re.search(r"[\[\]:*?/\\]", base) or base.startswith("'") or base.endswith("'"):
        raise ValueError("새 시트 이름은 1~31자이며 [ ] : * ? / \\ 와 맨 앞·뒤 작은따옴표를 사용할 수 없습니다.")
    title, suffix, used = base, 2, {s.casefold() for s in titles}
    while title.casefold() in used:
        tail = f" ({suffix})"
        title = base[:31-len(tail)] + tail
        suffix += 1
    return title


def _rows(workspace):
    active = workspace.active
    if not active:
        raise ValueError("저장할 옷이 없습니다.")
    invalid = [g.code for g in active if not g.export_ready]
    if invalid:
        raise ValueError("필수 입력·확인이 필요한 외투코드: " + ", ".join(invalid[:30]))
    rows = []
    for i, g in enumerate(active, 3):
        validate_fields(g.fields)
        rows.append([i-2, g.code, f'=IF(C{i}<>"", "0_" & C{i}, "")', g.fields["category"], "",
                     g.fields["gender"], g.fields["age"], g.fields["size"], ""] + [f'=IF(B{i}<>"",0,"")']*4)
    return rows


def _children(node, name):
    return [n for n in node.childNodes if n.nodeType == n.ELEMENT_NODE and n.localName == name]


def _first(node, name):
    return next(iter(_children(node, name)), None)


def _element(doc, parent, name, namespace=MAIN):
    prefix = parent.prefix
    return doc.createElementNS(namespace, f"{prefix}:{name}" if prefix else name)


def _package(raw):
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            names = archive.namelist()
            if len(names) != len(set(names)) or any(n.startswith("_xmlsignatures/") for n in names):
                raise ValueError("디지털 서명 또는 중복 항목이 있는 파일은 수정할 수 없습니다. 새 엑셀 파일에 저장하세요.")
            parts = {name: archive.read(name) for name in names}
        document = minidom.parseString(parts["xl/workbook.xml"])
        root = document.documentElement
        if root.namespaceURI != MAIN:
            raise ValueError("이 엑셀 파일 형식은 지원하지 않습니다. 표준 .xlsx 사본 또는 새 파일을 사용하세요.")
        protection = _first(root, "workbookProtection")
        if protection and protection.getAttribute("lockStructure") in {"1", "true"}:
            raise ValueError("통합 문서 구조가 보호되어 있습니다. 새 엑셀 파일에 저장하세요.")
        for required in ("xl/styles.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
            if required not in parts:
                raise ValueError("이 엑셀 파일의 구성이 지원되지 않습니다. 새 엑셀 파일에 저장하세요.")
        titles = [s.getAttribute("name") for s in _children(_first(root, "sheets"), "sheet")]
        return parts, titles
    except (BadZipFile, KeyError, OSError, ExpatError, AttributeError):
        raise ValueError("엑셀 파일을 읽을 수 없습니다. 암호화되지 않은 정상 .xlsx 파일을 선택하세요.") from None


def _target(path):
    target = Path(path).expanduser().resolve()
    if target.suffix.lower() != ".xlsx" or not target.name[:-5]:
        raise ValueError("저장 파일은 .xlsx 형식이어야 합니다.")
    if not target.parent.is_dir():
        raise ValueError("저장할 폴더가 없습니다. 다른 폴더를 선택하세요.")
    if target.is_dir():
        raise ValueError("폴더가 아닌 엑셀 파일 경로를 선택하세요.")
    return target


def _check_owner_file(path):
    if path.with_name("~$" + path.name).exists():
        raise ValueError("엑셀에서 사용 중인 파일입니다. 파일을 닫은 뒤 다시 저장하세요.")


def build_plan(workspace: Workspace, path: str | Path, base_title: str) -> ExportPlan:
    target = _target(path)
    _check_owner_file(target)
    rows = _rows(workspace)
    raw = target.read_bytes() if target.exists() else None
    titles = _package(raw)[1] if raw is not None else []
    return ExportPlan(workspace.id, workspace.export_fingerprint(), str(target),
                      hashlib.sha256(raw).hexdigest() if raw is not None else None,
                      next_title(base_title, titles), rows, workspace.counts, fingerprint(_template()))


def _rgb(value):
    channels = [round(max(0, min(1, value.get(key, 0))) * 255) for key in ("red", "green", "blue")]
    return "FF" + "".join(f"{channel:02X}" for channel in channels)


def _style(cell, value):
    text = value.get("textFormat", {})
    cell.font = Font(name=text.get("fontFamily", "Malgun Gothic"), size=text.get("fontSize", 11),
                     bold=text.get("bold", False), italic=text.get("italic", False),
                     color=_rgb(text.get("foregroundColor", {})))
    if "backgroundColor" in value:
        cell.fill = PatternFill("solid", fgColor=_rgb(value["backgroundColor"]))
    cell.alignment = Alignment(horizontal=value.get("horizontalAlignment", "general").lower(),
                               vertical={"MIDDLE": "center"}.get(value.get("verticalAlignment"),
                                   value.get("verticalAlignment", "BOTTOM").lower()),
                               wrap_text=value.get("wrapStrategy") == "WRAP")
    borders = {}
    for side, border in value.get("borders", {}).items():
        borders[side] = Side(style={"SOLID": "thin", "SOLID_MEDIUM": "medium", "SOLID_THICK": "thick",
                                    "DASHED": "dashed", "DOTTED": "dotted", "DOUBLE": "double"}.get(border.get("style"), "thin"),
                             color=_rgb(border.get("color", {})))
    cell.border = Border(**borders)
    cell.number_format = value.get("numberFormat", {}).get("pattern", "General")


def _new_workbook(plan, data):
    book = Workbook()
    sheet = book.active
    sheet.title = plan.title
    total = max(len(plan.rows)+2, data["row_count"])
    for r, c, item in data["cells"]:
        cell = sheet.cell(r+1, c+1)
        _style(cell, item.get("userEnteredFormat", {}))
        value = item.get("userEnteredValue", {})
        if value:
            cell.value = next(iter(value.values()))
            if "stringValue" in value:
                cell.data_type = "s"
        validation = item.get("dataValidation", {})
        if validation:
            choices = [str(v["userEnteredValue"]) for v in validation["condition"]["values"]]
            rule = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True,
                                  showErrorMessage=bool(validation.get("strict")), showDropDown=False)
            sheet.add_data_validation(rule)
            rule.add(f"{get_column_letter(c+1)}3:{get_column_letter(c+1)}{total}")
    for r in range(4, total+1):
        for c in range(2, 15):
            sheet.cell(r, c)._style = copy.copy(sheet.cell(3, c)._style)
    for r in range(1, total+1):
        sheet.row_dimensions[r].height = data["row_heights"][min(r-1, 2)] * .75
    for c, pixels in enumerate(data["column_widths"], 1):
        sheet.column_dimensions[get_column_letter(c)].width = max(0, (pixels-5)/7)
    for r, values in enumerate(plan.rows, 3):
        for c, value in enumerate(values, 2):
            cell = sheet.cell(r, c, value)
            if c not in (2, 4, 11, 12, 13, 14):
                cell.data_type = "s"
        sheet.cell(r, 3).number_format = "@"
    sheet.freeze_panes = "E3"
    output = io.BytesIO()
    book.save(output)
    book.close()
    # Supply formula caches for viewers which do not calculate Excel formulas.
    with ZipFile(io.BytesIO(output.getvalue())) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    document = minidom.parseString(parts["xl/worksheets/sheet1.xml"])
    for cell in document.getElementsByTagNameNS(MAIN, "c"):
        if not _first(cell, "f"):
            continue
        address = cell.getAttribute("r")
        if address == "C1":
            cached = "일일수량: " + str(len(plan.rows))
        elif address.startswith("D"):
            index = int(address[1:])-3
            cached = "0_" + plan.rows[index][1] if 0 <= index < len(plan.rows) else ""
        else:
            cached = 0
        value = _first(cell, "v") or _element(document, cell, "v")
        while value.firstChild:
            value.removeChild(value.firstChild)
        value.appendChild(document.createTextNode(str(cached)))
        if value.parentNode is None:
            cell.appendChild(value)
        if isinstance(cached, str):
            cell.setAttribute("t", "str")
    parts["xl/worksheets/sheet1.xml"] = document.toxml(encoding="utf-8")
    return parts


def _merge_styles(original, incoming):
    target = minidom.parseString(original)
    source = minidom.parseString(incoming)
    root = target.documentElement
    maps = {}
    order = ("numFmts", "fonts", "fills", "borders", "cellStyleXfs", "cellXfs", "cellStyles",
             "dxfs", "tableStyles", "colors", "extLst")
    for group, item_name in (("numFmts", "numFmt"), ("fonts", "font"), ("fills", "fill"),
                             ("borders", "border"), ("cellStyleXfs", "xf"), ("cellXfs", "xf")):
        src = _first(source.documentElement, group)
        dst = _first(root, group)
        if src is None:
            maps[group] = {}
            continue
        if dst is None:
            dst = _element(target, root, group)
            following = next((n for n in root.childNodes if n.nodeType == n.ELEMENT_NODE and
                              (order.index(n.localName) if n.localName in order else len(order)) > order.index(group)), None)
            root.insertBefore(dst, following)
        existing = _children(dst, item_name)
        remap = {}
        next_num_format = max([163] + [int(v.getAttribute("numFmtId")) for v in existing
                                      if v.hasAttribute("numFmtId")])+1
        for index, item in enumerate(_children(src, item_name)):
            node = target.importNode(item, True)
            # Imported nodes use the incoming default namespace. A destination
            # with only an x:styleSheet prefix has no inherited default, so keep
            # the declaration on each imported subtree when serializing it.
            node.setAttributeNS("http://www.w3.org/2000/xmlns/", "xmlns", MAIN)
            if group == "numFmts":
                remap[int(item.getAttribute("numFmtId"))] = next_num_format
                node.setAttribute("numFmtId", str(next_num_format))
                next_num_format += 1
            else:
                remap[index] = len(existing)
                for key, owner in (("fontId", "fonts"), ("fillId", "fills"), ("borderId", "borders"),
                                   ("numFmtId", "numFmts"), ("xfId", "cellStyleXfs")):
                    if node.hasAttribute(key):
                        old = int(node.getAttribute(key))
                        node.setAttribute(key, str(maps.get(owner, {}).get(old, old)))
            dst.appendChild(node)
            existing.append(node)
        dst.setAttribute("count", str(len(existing)))
        maps[group] = remap
    return target.toxml(encoding="utf-8"), maps["cellXfs"]


def _append(original, incoming, title):
    parts = original.copy()
    styles, mapping = _merge_styles(parts["xl/styles.xml"], incoming["xl/styles.xml"])
    parts["xl/styles.xml"] = styles
    sheet = minidom.parseString(incoming["xl/worksheets/sheet1.xml"])
    for node in sheet.getElementsByTagName("*"):
        for attribute in ("s", "style"):
            if node.hasAttribute(attribute):
                node.setAttribute(attribute, str(mapping[int(node.getAttribute(attribute))]))
    number = 1
    while f"xl/worksheets/sheet{number}.xml" in parts:
        number += 1
    sheet_path = f"xl/worksheets/sheet{number}.xml"
    parts[sheet_path] = sheet.toxml(encoding="utf-8")
    relationships = minidom.parseString(parts["xl/_rels/workbook.xml.rels"])
    ids = {n.getAttribute("Id") for n in relationships.getElementsByTagNameNS(REL, "Relationship")}
    index = 1
    while f"rId{index}" in ids:
        index += 1
    relation = _element(relationships, relationships.documentElement, "Relationship", REL)
    for key, value in {"Id": f"rId{index}", "Type": DOC_REL+"/worksheet", "Target": "worksheets/"+Path(sheet_path).name}.items():
        relation.setAttribute(key, value)
    relationships.documentElement.appendChild(relation)
    parts["xl/_rels/workbook.xml.rels"] = relationships.toxml(encoding="utf-8")
    workbook = minidom.parseString(parts["xl/workbook.xml"])
    sheets = _first(workbook.documentElement, "sheets")
    node = _element(workbook, sheets, "sheet")
    node.setAttribute("name", title)
    node.setAttribute("sheetId", str(max(int(n.getAttribute("sheetId")) for n in _children(sheets, "sheet"))+1))
    node.setAttributeNS("http://www.w3.org/2000/xmlns/", "xmlns:r", DOC_REL)
    node.setAttributeNS(DOC_REL, "r:id", f"rId{index}")
    sheets.appendChild(node)
    parts["xl/workbook.xml"] = workbook.toxml(encoding="utf-8")
    content = minidom.parseString(parts["[Content_Types].xml"])
    override = _element(content, content.documentElement, "Override", CONTENT)
    override.setAttribute("PartName", "/"+sheet_path)
    override.setAttribute("ContentType", "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml")
    content.documentElement.appendChild(override)
    parts["[Content_Types].xml"] = content.toxml(encoding="utf-8")
    return parts


@contextmanager
def _read_guard(path):
    """Windows denies writers while the original is checked and replaced."""
    handle = None
    if sys.platform == "win32" and path.exists():
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.CreateFileW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p,
                                      ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p]
        kernel.CreateFileW.restype = ctypes.c_void_p
        kernel.CloseHandle.argtypes = [ctypes.c_void_p]
        handle = kernel.CreateFileW(str(path), 0x80000000, 0x00000001 | 0x00000004, None, 3, 0x80, None)
        if handle == ctypes.c_void_p(-1).value:
            raise ValueError("엑셀 파일이 사용 중이거나 읽을 수 없습니다. 파일을 닫은 뒤 다시 저장하세요.")
    try:
        yield
    finally:
        if handle is not None:
            kernel.CloseHandle(handle)


def _digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None


def _replace(source, target):
    if sys.platform == "win32":
        # Unlike MoveFileEx (os.replace), ReplaceFile works while our read guard
        # denies other writers and permits deletion of the destination handle.
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.ReplaceFileW.argtypes = [ctypes.c_wchar_p, ctypes.c_wchar_p, ctypes.c_wchar_p,
                                       ctypes.c_uint32, ctypes.c_void_p, ctypes.c_void_p]
        kernel.ReplaceFileW.restype = ctypes.c_int
        if not kernel.ReplaceFileW(str(target), str(source), None, 0, None, None):
            raise ctypes.WinError(ctypes.get_last_error())
    else:
        os.replace(source, target)


def save(repository: Repository, plan: ExportPlan) -> str:
    workspace = repository.load(plan.workspace_id)
    if workspace.export_fingerprint() != plan.workspace_digest or _rows(workspace) != plan.rows:
        raise ValueError("작업 내용이 변경되었습니다. 미리보기를 다시 만드세요.")
    data = _template()
    if fingerprint(data) != plan.template_digest:
        raise ValueError("내장 양식이 변경되었습니다. 미리보기를 다시 만드세요.")
    target = _target(plan.path)
    _check_owner_file(target)
    lock = QLockFile(str(target)+".coatdesk.lock")
    if not lock.tryLock(0):
        raise ValueError("다른 작업에서 이 엑셀 파일에 저장 중입니다. 잠시 후 다시 시도하세요.")
    temporary = None
    try:
        with _read_guard(target):
            if _digest(target) != plan.file_digest:
                raise ValueError("미리보기 후 엑셀 파일이 변경되었습니다. 미리보기를 다시 만드세요.")
            original = _package(target.read_bytes())[0] if target.exists() else None
            if original:
                if next_title(plan.title, _package(target.read_bytes())[1]) != plan.title:
                    raise ValueError("같은 이름의 시트가 생겼습니다. 미리보기를 다시 만드세요.")
            else:
                next_title(plan.title, [])
            parts = _new_workbook(plan, data)
            if original:
                parts = _append(original, parts, plan.title)
                if any(parts[name] != raw for name, raw in original.items() if name not in MODIFIED_PARTS):
                    raise ValueError("기존 시트 보존 검사에 실패했습니다. 원본 파일은 변경하지 않았습니다.")
            descriptor, name = tempfile.mkstemp(prefix=".coatdesk-", suffix=".xlsx", dir=target.parent)
            temporary = Path(name)
            with os.fdopen(descriptor, "wb") as stream:
                with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
                    for name, raw in parts.items():
                        archive.writestr(name, raw)
                stream.flush()
                os.fsync(stream.fileno())
            book = load_workbook(temporary, read_only=True, data_only=False, keep_links=True)
            try:
                sheet = book[plan.title]
                for r, values in enumerate(plan.rows, 3):
                    if [sheet.cell(r, c).value or "" for c in range(2, 15)] != values:
                        raise ValueError("저장 결과 확인에 실패했습니다. 원본 파일은 변경하지 않았습니다.")
            finally:
                book.close()
            if _digest(target) != plan.file_digest:
                raise ValueError("저장 중 엑셀 파일이 변경되었습니다. 미리보기를 다시 만드세요.")
            if plan.file_digest is None:
                if sys.platform == "win32":
                    os.rename(temporary, target)
                else:
                    os.link(temporary, target)
                    temporary.unlink()
            else:
                _replace(temporary, target)
            temporary = None
        return str(target)
    except (OSError, BadZipFile) as error:
        raise ValueError("엑셀을 저장하지 못했습니다. 파일을 닫고 저장 폴더·권한·여유 공간을 확인하세요.") from error
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        lock.unlock()
