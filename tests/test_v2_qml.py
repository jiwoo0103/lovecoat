import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("QT_QUICK_BACKEND", "software")

from PySide6.QtCore import Q_ARG, QMetaObject, QObject, QUrl, QPoint, QPointF, Qt
from PySide6.QtGui import QFontDatabase
from PySide6.QtQml import QQmlApplicationEngine
from PySide6.QtQuickControls2 import QQuickStyle
from PySide6.QtTest import QTest

from coatdesk_v2.backend import Backend
from coatdesk_v2.images import Photos
from coatdesk_v2.scanner import scan
from test_v2_workspace import application, photo
from test_v2_workspace import wait_for
from openpyxl import load_workbook


def test_qml_review_photo_zoom_dialogs_and_small_screen_have_no_runtime_warnings(tmp_path):
    app = application()
    QQuickStyle.setStyle("Basic")
    for code in ("0001", "0002", "0003", "0004", "0005"):
        for angle in range(3):
            photo(tmp_path / "photos", code, angle)
    backend = Backend(tmp_path / "data")
    key = backend.repository.create("UI 검증", "", tmp_path / "photos", 1, 5, "", scan(tmp_path / "photos", 1, 5))
    backend.openWorkspace(key)
    engine = QQmlApplicationEngine()
    engine.rootContext().setContextProperty("appBackend", backend)
    engine.addImageProvider("photos", Photos())
    warnings = []
    engine.warnings.connect(lambda messages: warnings.extend(m.toString() for m in messages))
    engine.load(QUrl.fromLocalFile(str(Path(__file__).parents[1] / "coatdesk_v2/qml/Main.qml")))
    assert engine.rootObjects(), warnings
    window = engine.rootObjects()[0]
    QTest.qWait(250)
    backend.selectPage()
    backend.batchFields("코트", "여성", "성인", "M")
    backend.confirmSelected()
    assert backend.state["counts"]["confirmed"] == 4
    QTest.qWait(50)
    def find_item(item, name):
        if item.objectName() == name:
            return item
        for child in item.childItems():
            result = find_item(child, name)
            if result is not None:
                return result
        return None
    picture = find_item(window.contentItem(), "photo_0001")
    assert picture is not None
    point = picture.mapToScene(QPointF(picture.width()/2, picture.height()/2))
    QTest.mouseClick(window, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, QPoint(int(point.x()), int(point.y())))
    QTest.keyClick(window, Qt.Key.Key_Right)
    assert backend.page == 1
    for mode, front in (("review", True), ("photos", True), ("photos", False)):
        backend.setMode(mode)
        backend.viewOptions(4, 2, front, 150)
        window.resize(1000, 720)
        QTest.qWait(100)
        assert not window.grabWindow().isNull()
    for name in ("createDialog", "sourceDialog", "batchDialog", "viewDialog", "headDialog", "zoomDialog", "transferDialog"):
        dialog = window.findChild(QObject, name)
        assert dialog is not None
        QMetaObject.invokeMethod(dialog, "open")
        QTest.qWait(30)
        QMetaObject.invokeMethod(dialog, "close")
    QMetaObject.invokeMethod(window, "showZoom", Q_ARG("QVariant", "0001"))
    QTest.qWait(80)
    backend.nextGarment(1)
    QTest.qWait(80)
    assert backend.current == "0002"
    assert not warnings, warnings
    window.close()
    engine.deleteLater()
    QTest.qWait(50)
    backend.shutdown()
    assert app


def test_qml_excel_dialog_appends_two_dates_and_preserves_first_sheet(tmp_path):
    app = application()
    font = Path("C:/Windows/Fonts/malgun.ttf")
    if font.exists():
        QFontDatabase.addApplicationFont(str(font))
    QQuickStyle.setStyle("Basic")
    photo(tmp_path / "photos", "0001")
    backend = Backend(tmp_path / "data")
    keys = [backend.repository.create(day, "", tmp_path / "photos", 1, 1, "", scan(tmp_path / "photos", 1, 1))
            for day in ("20261002", "20261003")]
    for key in keys:
        backend.openWorkspace(key)
        backend.selectPage()
        backend.batchFields("자켓", "여성", "성인", "M")
    engine = QQmlApplicationEngine()
    engine.rootContext().setContextProperty("appBackend", backend)
    engine.addImageProvider("photos", Photos())
    warnings = []
    engine.warnings.connect(lambda messages: warnings.extend(m.toString() for m in messages))
    engine.load(QUrl.fromLocalFile(str(Path(__file__).parents[1] / "coatdesk_v2/qml/Main.qml")))
    assert engine.rootObjects(), warnings
    window = engine.rootObjects()[0]
    window.resize(1000, 720)
    path = tmp_path / "외투 목록.xlsx"
    dialog = window.findChild(QObject, "transferDialog")
    def click_named(name):
        item = window.findChild(QObject, name)
        assert item is not None and item.property("enabled"), name
        point = item.mapToScene(QPointF(item.width()/2, item.height()/2))
        assert 0 <= point.y() < window.height(), name
        QTest.mouseClick(window, Qt.LeftButton, Qt.NoModifier, QPoint(round(point.x()), round(point.y())))
        wait_for(lambda: not backend.jobs, 15000)
        QTest.qWait(80)
    for key, day in zip(keys, ("20261002", "20261003"), strict=True):
        backend.openWorkspace(key)
        QMetaObject.invokeMethod(dialog, "open")
        QTest.qWait(80)
        window.findChild(QObject, "excelPathField").setProperty("text", str(path))
        assert window.findChild(QObject, "newExcelSheetName").property("text") == day
        click_named("excelPreviewButton")
        assert backend.plan is not None, backend.error
        assert backend.plan.title == day
        assert backend.plan.path == str(path)
        capture_folder = os.environ.get("COATDESK_UI_SCREENSHOTS")
        if capture_folder:
            Path(capture_folder).mkdir(parents=True, exist_ok=True)
            window.grabWindow().save(str(Path(capture_folder) / f"excel-preview-{day}-1000.png"))
        click_named("excelSaveButton")
        assert not backend.error, backend.error
        assert backend.last_path == str(path)
        QMetaObject.invokeMethod(dialog, "close")
        QTest.qWait(40)
    with path.open("rb") as stream:
        book = load_workbook(stream)
        assert book.sheetnames == ["20261002", "20261003"]
        for sheet in book:
            assert sheet["C3"].value == "0001"
            assert sheet["E3"].value == "자켓"
            assert sheet["G3"].value == "여성"
        book.close()
    # The most recently chosen workbook is remembered after a restart.
    assert backend.repository.setting("excel_path") == str(path)
    assert not warnings, warnings
    window.close()
    engine.deleteLater()
    QTest.qWait(80)
    backend.shutdown()
    assert app
