import copy
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.packaging.relationship import Relationship
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.external_link.external import ExternalBook, ExternalLink
from PySide6.QtCore import QLockFile

from coatdesk_v2 import excel
from coatdesk_v2.domain import Garment
from coatdesk_v2.storage import Repository


@pytest.fixture
def prepared(tmp_path):
    repo = Repository(tmp_path / "data")
    garments = [Garment(f"{i:04}", {}, fields={"category": "자켓", "gender": "여성", "age": "성인", "size": "M"})
                for i in range(1, 5)]
    key = repo.create("20261002", "", tmp_path / "photos", 1, 4, "2", garments)
    return repo, key, tmp_path / "외투목록.xlsx"


def package(path):
    with ZipFile(path) as archive:
        return {n: archive.read(n) for n in archive.namelist()}


def test_new_workbook_then_second_date_and_duplicate_keep_existing_data(prepared):
    repo, key, target = prepared
    plan = excel.build_plan(repo.load(key), target, "20261002")
    assert plan.path == str(target.resolve())
    assert plan.counts == {"total": 3, "confirmed": 3, "excluded": 1, "missing": 0}
    assert excel.save(repo, plan) == str(target.resolve())
    before = package(target)
    plan = excel.build_plan(repo.load(key), target, "20261003")
    excel.save(repo, plan)
    after = package(target)
    assert all(after[n] == raw for n, raw in before.items() if n not in excel.MODIFIED_PARTS)
    plan = excel.build_plan(repo.load(key), target, "20261003")
    assert plan.title == "20261003 (2)"
    excel.save(repo, plan)
    book = load_workbook(target)
    try:
        assert book.sheetnames == ["20261002", "20261003", "20261003 (2)"]
        for name in book.sheetnames:
            sheet = book[name]
            assert [sheet.cell(r, 2).value for r in range(3, 6)] == [1, 2, 3]
            assert [sheet.cell(r, 3).value for r in range(3, 6)] == ["0001", "0003", "0004"]
            assert sheet["C3"].number_format == "@"
            assert sheet["E3"].value == "자켓"
            assert sheet["F3"].value is None and sheet["J3"].value is None
            assert sheet["D5"].value == '=IF(C5<>"", "0_" & C5, "")'
            assert sheet["K5"].value == '=IF(B5<>"",0,"")'
            assert sheet["B6"].value is None
            assert sheet["B2"].fill.fgColor.rgb == "FFD9E2F3"
            assert sheet["B3"].border.top.style == "thin"
            assert sheet.row_dimensions[1].height == 79.5
            assert sheet.row_dimensions[3].height == 17.25
            assert sheet.column_dimensions["C"].width == pytest.approx(20)
            assert sheet["P3"].value == "자켓"
            assert sheet["Q3"].value == "2XS"
            assert sheet["R3"].value == "60/70"
            assert {str(rule.sqref) for rule in sheet.data_validations.dataValidation} == {"E3:E1000", "I3:I1000"}
    finally:
        book.close()
    book = load_workbook(target, data_only=True)
    try:
        assert book["20261003"]["C1"].value == "일일수량: 3"
        assert book["20261003"]["D5"].value == "0_0004"
        assert book["20261003"]["K5"].value == 0
    finally:
        book.close()


def test_existing_chart_external_links_styles_and_custom_package_parts_survive(prepared):
    repo, key, target = prepared
    book = Workbook()
    sheet = book.active
    sheet.title = "기존 업무"
    for row in (("분류", "수량"), ("원본", 10), ("원본2", 20)):
        sheet.append(row)
    sheet["A2"].font = Font(name="Arial", bold=True, color="AA5500")
    sheet["A2"].fill = PatternFill("solid", fgColor="AAFFCC")
    sheet["B2"].number_format = '#,##0.00"건"'
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    sheet.add_chart(chart, "D2")
    external = ExternalLink(externalBook=ExternalBook(id="rId1"))
    external.file_link = Relationship(Type=excel.DOC_REL+"/externalLinkPath", Target="file:///original-source.xlsx",
                                      TargetMode="External")
    book._external_links.append(external)
    book.save(target)
    book.close()
    with ZipFile(target, "a") as archive:
        archive.writestr("customXml/custom-preserved.xml", '<sample xmlns="urn:user">keep all content</sample>')
    before = package(target)
    excel.save(repo, excel.build_plan(repo.load(key), target, "20261003"))
    after = package(target)
    assert all(after[n] == raw for n, raw in before.items() if n not in excel.MODIFIED_PARTS)
    assert any(n.startswith("xl/charts/") for n in after)
    assert any(n.startswith("xl/externalLinks/") for n in after)
    book = load_workbook(target)
    try:
        assert book["기존 업무"]["A2"].value == "원본"
        assert book["기존 업무"]["A2"].fill.fgColor.rgb == "00AAFFCC"
        assert book["기존 업무"]["B2"].number_format == '#,##0.00"건"'
        assert len(book["기존 업무"]._charts) == 1
        assert len(book._external_links) == 1
        assert book["20261003"]["E3"].value == "자켓"
        assert book["20261003"]["B2"].fill.fgColor.rgb == "FFD9E2F3"
    finally:
        book.close()


def test_existing_prefixed_package_namespaces_keep_new_styles_in_main_namespace(prepared):
    repo, key, target = prepared
    book = Workbook()
    book.active.title = "기존"
    book.active["A1"] = "기존 자료"
    book.active["A1"].font = Font(name="Arial", bold=True, color="112233")
    book.save(target)
    book.close()
    original = package(target)
    namespaces = {
        "xl/styles.xml": ("x", excel.MAIN),
        "xl/workbook.xml": ("x", excel.MAIN),
        "xl/_rels/workbook.xml.rels": ("rel", excel.REL),
        "[Content_Types].xml": ("ct", excel.CONTENT),
    }
    for name, (prefix, namespace) in namespaces.items():
        ET.register_namespace(prefix, namespace)
        original[name] = ET.tostring(ET.fromstring(original[name]), encoding="utf-8", xml_declaration=True)
        assert b'xmlns="' not in original[name]
    with ZipFile(target, "w") as archive:
        for name, raw in original.items():
            archive.writestr(name, raw)
    excel.save(repo, excel.build_plan(repo.load(key), target, "20261003"))
    appended = package(target)
    assert all(appended[name] == raw for name, raw in original.items() if name not in excel.MODIFIED_PARTS)
    styles = ET.fromstring(appended["xl/styles.xml"])
    for name in ("fonts", "fills", "borders", "cellStyleXfs", "cellXfs"):
        group = styles.find(f"{{{excel.MAIN}}}{name}")
        assert group is not None
        assert all(item.tag.startswith("{"+excel.MAIN+"}") for child in group for item in child.iter())
    workbook = ET.fromstring(appended["xl/workbook.xml"])
    sheets = workbook.find(f"{{{excel.MAIN}}}sheets")
    assert [s.attrib["name"] for s in sheets] == ["기존", "20261003"]
    assert all(s.tag == f"{{{excel.MAIN}}}sheet" for s in sheets)
    relationships = ET.fromstring(appended["xl/_rels/workbook.xml.rels"])
    assert all(r.tag == f"{{{excel.REL}}}Relationship" for r in relationships)
    content = ET.fromstring(appended["[Content_Types].xml"])
    assert all(n.tag in {f"{{{excel.CONTENT}}}Default", f"{{{excel.CONTENT}}}Override"} for n in content)
    book = load_workbook(target)
    try:
        assert book["기존"]["A1"].value == "기존 자료"
        assert book["기존"]["A1"].font.color.rgb == "00112233"
        assert book["20261003"]["E3"].value == "자켓"
        assert book["20261003"]["E3"].font.name == "Malgun Gothic"
        assert book["20261003"]["E3"].font.color.rgb == "FF000000"
        assert book["20261003"]["E2"].fill.fgColor.rgb == "FFD9E2F3"
    finally:
        book.close()


@pytest.mark.parametrize("title", ["", "x"*32, "a/b", "a\\b", "a[b", "a:b", "a*b", "a?b", "'x", "x'"])
def test_sheet_names_follow_excel_limits(title):
    with pytest.raises(ValueError):
        excel.next_title(title, [])


def test_duplicate_names_are_case_insensitive_and_suffix_fits_31_characters():
    assert excel.next_title("date", ["DATE"]) == "date (2)"
    base = "a"*31
    assert excel.next_title(base, [base]) == "a"*27+" (2)"


@pytest.mark.parametrize("kind", ["missing", "review", "empty"])
def test_incomplete_or_review_pending_or_empty_workspace_cannot_export(prepared, kind):
    repo, key, target = prepared
    if kind == "missing":
        repo.edit(key, ["0001"], lambda g: g.set_field("category", ""))
    elif kind == "review":
        repo.edit(key, ["0001"], lambda g: setattr(g, "review_required", ["size"]))
    else:
        repo.update(key, errors=[1, 2, 3, 4])
    with pytest.raises(ValueError):
        excel.build_plan(repo.load(key), target, "20261002")
    assert not target.exists()


def test_new_save_cannot_be_replayed_and_changed_file_or_work_is_blocked(prepared):
    repo, key, target = prepared
    plan = excel.build_plan(repo.load(key), target, "20261002")
    excel.save(repo, plan)
    original = target.read_bytes()
    with pytest.raises(ValueError, match="파일이 변경"):
        excel.save(repo, plan)
    assert target.read_bytes() == original
    plan = excel.build_plan(repo.load(key), target, "20261003")
    repo.edit(key, ["0001"], lambda g: g.set_field("size", "XL"))
    with pytest.raises(ValueError, match="작업 내용"):
        excel.save(repo, plan)
    assert target.read_bytes() == original
    plan = excel.build_plan(repo.load(key), target, "20261003")
    with ZipFile(target, "a") as archive:
        archive.writestr("customXml/changed.xml", "changed after preview")
    changed = target.read_bytes()
    with pytest.raises(ValueError, match="파일이 변경"):
        excel.save(repo, plan)
    assert target.read_bytes() == changed


def test_excel_owner_lock_and_other_app_writer_block_without_writing(prepared):
    repo, key, target = prepared
    plan = excel.build_plan(repo.load(key), target, "20261002")
    owner = target.with_name("~$"+target.name)
    owner.write_text("excel owner", encoding="utf-8")
    with pytest.raises(ValueError, match="사용 중"):
        excel.save(repo, plan)
    with pytest.raises(ValueError, match="사용 중"):
        excel.build_plan(repo.load(key), target, "20261002")
    owner.unlink()
    lock = QLockFile(str(target)+".coatdesk.lock")
    assert lock.tryLock(0)
    try:
        with pytest.raises(ValueError, match="저장 중"):
            excel.save(repo, plan)
    finally:
        lock.unlock()
    assert not target.exists()


def test_failed_replace_preserves_existing_file_and_removes_only_temporary_output(prepared, monkeypatch):
    repo, key, target = prepared
    excel.save(repo, excel.build_plan(repo.load(key), target, "20261002"))
    before = target.read_bytes()
    plan = excel.build_plan(repo.load(key), target, "20261003")
    def fail(*args):
        raise PermissionError("simulated locked workbook")
    monkeypatch.setattr(excel, "_replace", fail)
    with pytest.raises(ValueError, match="저장하지 못"):
        excel.save(repo, plan)
    assert target.read_bytes() == before
    assert not list(target.parent.glob(".coatdesk-*"))
    assert not Path(str(target)+".coatdesk.lock").exists()


def test_recheck_detects_invalid_candidate_before_original_is_replaced(prepared, monkeypatch):
    repo, key, target = prepared
    excel.save(repo, excel.build_plan(repo.load(key), target, "20261002"))
    before = target.read_bytes()
    original = excel._new_workbook
    def corrupt(*args):
        parts = original(*args)
        parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(b">0001<", b">9999<")
        return parts
    monkeypatch.setattr(excel, "_new_workbook", corrupt)
    with pytest.raises(ValueError, match="저장 결과 확인"):
        excel.save(repo, excel.build_plan(repo.load(key), target, "20261003"))
    assert target.read_bytes() == before


def test_user_size_is_written_as_text_and_never_executed_as_a_formula(prepared):
    repo, key, target = prepared
    value = '=HYPERLINK("https://example.invalid","text")'
    repo.edit(key, ["0001"], lambda g: g.set_field("size", value))
    excel.save(repo, excel.build_plan(repo.load(key), target, "20261002"))
    book = load_workbook(target)
    try:
        assert book["20261002"]["I3"].value == value
        assert book["20261002"]["I3"].data_type == "s"
    finally:
        book.close()


def test_unsupported_corrupt_signed_and_structure_locked_files_are_not_modified(prepared):
    repo, key, target = prepared
    target.write_bytes(b"not a spreadsheet")
    with pytest.raises(ValueError):
        excel.build_plan(repo.load(key), target, "20261002")
    assert target.read_bytes() == b"not a spreadsheet"
    book = Workbook()
    book.security.lockStructure = True
    book.save(target)
    book.close()
    protected = target.read_bytes()
    with pytest.raises(ValueError, match="구조가 보호"):
        excel.build_plan(repo.load(key), target, "20261002")
    assert target.read_bytes() == protected
    book = Workbook()
    book.save(target)
    book.close()
    with ZipFile(target, "a") as archive:
        archive.writestr("_xmlsignatures/sig1.xml", "signature")
    signed = target.read_bytes()
    with pytest.raises(ValueError, match="디지털 서명"):
        excel.build_plan(repo.load(key), target, "20261002")
    assert target.read_bytes() == signed


def test_plan_rows_cannot_be_changed_after_preview(prepared):
    repo, key, target = prepared
    plan = excel.build_plan(repo.load(key), target, "20261002")
    changed = copy.deepcopy(plan)
    changed.rows[0][3] = "코트"
    with pytest.raises(ValueError, match="작업 내용"):
        excel.save(repo, changed)
    assert not target.exists()
